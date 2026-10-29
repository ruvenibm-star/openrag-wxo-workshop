#!/usr/bin/env -S uv run --script
# /// script
# requires-python = ">=3.11"
# dependencies = ["reportlab>=4", "python-docx>=1.1", "python-pptx>=1.0", "pillow>=10"]
# ///
"""Build the workshop's sample documents for Blue Defence, a fictional Singapore defence company.

Every fact the labs ask about is written here, and answer-key.yaml quotes the same values.
Change both together.

    scripts/build-sample-docs.py            # writes sample-docs/
    scripts/build-sample-docs.py --out DIR

Output:
    blue-defence-knowledge-base/   the folder you upload in Lab 1 (5 files, 4 formats)
    scanned/                       an image-only PDF for the OCR exercise
"""

import argparse
import csv
import datetime as dt
import os
import random
from pathlib import Path

import reportlab
from docx import Document
from docx.shared import Pt
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from pptx import Presentation
from pptx.util import Inches
from pptx.util import Pt as PptPt
from reportlab import rl_config
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

AGENCY = "Blue Defence"
BANNER = "FICTIONAL SAMPLE DOCUMENT FOR WORKSHOP USE. Blue Defence and the content of this document are fictional."
FIXED_TIME = dt.datetime(2026, 10, 1, 9, 0, 0)

# ---------------------------------------------------------------------------------------------
# Document content
# ---------------------------------------------------------------------------------------------

SECURITY = {
    2024: {
        "ref": "POL-SEC-011 Revision 1",
        "meta": "Effective 1 January 2024. Owner: Security Office.",
        "sponsorship": None,
        "validity_text": "A clearance is valid for the period below from the date it is granted.",
        "levels": [
            "Staff who handle CONFIDENTIAL material need a CONFIDENTIAL clearance.",
            "Staff who handle SECRET or TOP SECRET material need a clearance at that level.",
        ],
        "validity": {"CONFIDENTIAL": ("10 years", "3 months"), "SECRET": ("7 years", "3 months"),
                     "TOP SECRET": ("5 years", "6 months")},
        "incidents": "Report a security incident to the Security Office within 72 hours of discovery.",
        "extra": [],
    },
    2026: {
        "ref": "POL-SEC-011 Revision 2",
        "meta": "Effective 1 July 2026. Owner: Security Office.",
        "sponsorship": "The applicant's Division Director sponsors each clearance application. The Security Office "
                       "starts vetting within 10 working days of receiving it.",
        "validity_text": "A clearance is valid for the period below from the date it is granted. Submit the renewal "
                         "application to the Security Office before the renewal date; a clearance that lapses is "
                         "suspended until it is renewed.",
        "levels": [
            "Staff who handle CONFIDENTIAL material need a CONFIDENTIAL clearance.",
            "Staff who handle SECRET or TOP SECRET material need a clearance at that level.",
            "Staff with remote access to the restricted network need at least a SECRET clearance.",
        ],
        "validity": {"CONFIDENTIAL": ("8 years", "3 months"), "SECRET": ("5 years", "6 months"),
                     "TOP SECRET": ("3 years", "6 months")},
        "incidents": "Report a security incident to the Security Office within 24 hours of discovery.",
        "extra": [
            ("6. Overseas travel", [
                "Staff with a SECRET or TOP SECRET clearance declare overseas travel to the Security Office on form "
                "SEC-T3 at least 10 working days before departure.",
            ]),
        ],
    },
}

MANUAL = {
    "title": "Sentinel-7 Unmanned Surface Vessel: Maintenance Manual",
    "ref": "MM-S7-001 Revision C",
    "meta": "Issued 2 February 2026 by the Maritime Systems Division.",
    "specs": [
        ["Parameter", "Value"],
        ["Length overall", "7.2 m"],
        ["Maximum speed", "25 knots"],
        ["Endurance", "72 hours at 8 knots"],
        ["Energy storage", "4 lithium iron phosphate battery modules, 48 kWh each"],
        ["Propulsion", "2 electric propulsion pods"],
    ],
    "schedule": [
        ["Component", "Task", "Interval"],
        ["Propulsion pod seals", "Inspect; replace if worn", "Every 250 operating hours"],
        ["Hull anodes", "Replace", "Every 6 months, or when 50% consumed"],
        ["Battery modules", "Capacity test", "Every 500 operating hours"],
        ["Radar mast hinge", "Lubricate", "Every 100 operating hours"],
        ["Satellite antenna cabling", "Inspect for chafing", "Every 3 months"],
    ],
    "torque": [
        ["Fastener", "Thread", "Torque"],
        ["Hull access panel bolts", "M8", "24 N·m"],
        ["Propulsion pod mounting bolts", "M12", "85 N·m"],
        ["Antenna base bolts", "M6", "10 N·m"],
    ],
    "faults": [
        ["Code", "Meaning", "Operator action"],
        ["E-104", "Propulsion pod overtemperature", "Reduce thrust to 40% and return to base"],
        ["E-221", "GNSS signal lost", "Switch to inertial navigation and alert the shore operator"],
        ["E-305", "Battery cell imbalance above 80 mV", "Stop charging and run the balancing routine"],
    ],
    "battery": [
        "Charge battery modules only when the ambient temperature is between 5 °C and 40 °C.",
        "If a battery module shows smoke, swelling or a temperature above 60 °C, isolate shore power, "
        "evacuate everyone to at least 15 m from the vessel, and call the fire service.",
    ],
}

INSPECTION_MEMO = [
    "BLUE DEFENCE",
    "FACILITIES SAFETY OFFICE",
    "",
    "MEMO  INS-2026-031",
    "",
    "Date:      14 August 2026",
    "To:        Head, Kestrel Point Test Facility",
    "From:      Facilities Safety Office",
    "Subject:   Site safety inspection, Kestrel Point",
    "",
    "Findings",
    "1. Three fire extinguishers in Hangar 2 were past",
    "   their service date.",
    "2. The battery storage room had no hazard signage.",
    "3. The eyewash station in Workshop B was blocked",
    "   by stored equipment.",
    "",
    "Corrective actions",
    "a. Replace every fire extinguisher in Hangar 2",
    "   by 31 August 2026.",
    "b. Install hazard signage at the battery storage",
    "   room by 21 August 2026.",
    "c. Clear access to the Workshop B eyewash station",
    "   immediately.",
    "",
    "A follow-up inspection is scheduled for",
    "15 September 2026.",
    "",
    "Signed: R. Menon, Facilities Safety Officer",
    "",
    "FICTIONAL SAMPLE DOCUMENT FOR WORKSHOP USE",
]

PROJECTS = [
    {
        "name": "Project HALCYON: Sentinel-7 fleet expansion",
        "bullets": [
            "Status: AMBER",
            "8 of 12 vessels delivered",
            "Cause: late delivery of battery modules from Volta Marine Energy",
            "Recovery plan: remaining 4 vessels delivered by 15 December 2026",
            "Budget SGD 18.6 million; SGD 11.2 million spent",
        ],
    },
    {
        "name": "Project LANTERN: classified document search pilot",
        "bullets": [
            "Status: GREEN",
            "240 pilot users across 6 divisions",
            "Go-live: 15 January 2027",
            "Next step: accessibility testing in November 2026",
        ],
    },
    {
        "name": "Project ORCHID: secure data centre migration",
        "bullets": [
            "Status: RED",
            "61 of 140 workloads migrated",
            "12% over budget (SGD 1.4 million)",
            "Decision needed: steering committee on 12 November 2026",
        ],
    },
]

SUPPLIERS = [
    ["supplier_id", "supplier", "category", "contract_no", "contract_expiry", "contact_email"],
    ["SUP-014", "Volta Marine Energy Pte Ltd", "Lithium battery modules", "BD/C/2024/118", "2027-06-30", "sales@volta-marine.example"],
    ["SUP-022", "Halyard Composites Pte Ltd", "Hull fabrication", "BD/C/2023/077", "2026-12-31", "projects@halyard.example"],
    ["SUP-031", "Bluefin Sensors Ltd", "Radar and sonar units", "BD/C/2025/009", "2028-03-31", "contracts@bluefin.example"],
    ["SUP-037", "Meridian Satcom Pte Ltd", "Satellite communications", "BD/C/2025/041", "2027-09-30", "support@meridian-satcom.example"],
    ["SUP-045", "Ironbark Fasteners", "Marine-grade fasteners", "BD/C/2024/203", "2026-11-30", "orders@ironbark.example"],
    ["SUP-052", "Northlight Data Centres", "Colocation services", "BD/C/2022/150", "2027-01-31", "accounts@northlight.example"],
    ["SUP-058", "Tidewater Training Services", "Operator training", "BD/C/2025/066", "2026-10-31", "bookings@tidewater.example"],
    ["SUP-063", "Copperleaf Electrical", "Shore power installation", "BD/C/2023/190", "2027-04-30", "jobs@copperleaf.example"],
]

# ---------------------------------------------------------------------------------------------
# Writers
# ---------------------------------------------------------------------------------------------

rl_config.invariant = 1  # fixed PDF timestamps and IDs, so a rebuild without content changes is byte-identical
STYLES = getSampleStyleSheet()
TABLE_STYLE = TableStyle([
    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f3a5f")),
    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
    ("FONTSIZE", (0, 0), (-1, -1), 9),
    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
    ("VALIGN", (0, 0), (-1, -1), "TOP"),
])


def pdf_table(rows, widths):
    cell = STYLES["BodyText"].clone("cell", fontSize=9, leading=11)
    head = cell.clone("head", textColor=colors.white, fontName="Helvetica-Bold")
    data = [[Paragraph(str(v), head if r == 0 else cell) for v in row] for r, row in enumerate(rows)]
    table = Table(data, colWidths=[w * mm for w in widths], repeatRows=1)
    table.setStyle(TABLE_STYLE)
    return table


def pdf_header(story, title, ref, meta):
    story += [
        Paragraph(BANNER, STYLES["Italic"]),
        Spacer(1, 4 * mm),
        Paragraph(AGENCY, STYLES["Heading3"]),
        Paragraph(title, STYLES["Title"]),
        Paragraph(f"{ref}. {meta}", STYLES["BodyText"]),
        Spacer(1, 4 * mm),
    ]


def build_pdf(path, story, title):
    doc = SimpleDocTemplate(str(path), pagesize=A4, title=title, author=AGENCY,
                            leftMargin=20 * mm, rightMargin=20 * mm, topMargin=18 * mm, bottomMargin=18 * mm)

    def footer(canvas, d):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.drawString(20 * mm, 10 * mm, f"{AGENCY} · {title} · page {d.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)


def write_manual(path):
    m = MANUAL
    story = []
    pdf_header(story, m["title"], m["ref"], m["meta"])
    story += [
        Paragraph("1. Vessel overview", STYLES["Heading2"]),
        Paragraph("Sentinel-7 is a 7.2 m electric unmanned surface vessel for harbour patrol and mine "
                  "countermeasure trials. This manual covers scheduled maintenance, fastener torques, fault codes "
                  "and battery safety.", STYLES["BodyText"]),
        Spacer(1, 2 * mm), pdf_table(m["specs"], [55, 110]),
        Paragraph("2. Scheduled maintenance", STYLES["Heading2"]),
        Paragraph("Record every task in the vessel logbook. Operating hours are read from the vessel "
                  "management computer.", STYLES["BodyText"]),
        Spacer(1, 2 * mm), pdf_table(m["schedule"], [50, 55, 60]),
        PageBreak(),
        Paragraph("3. Fastener torques", STYLES["Heading2"]),
        Paragraph("Use a calibrated torque wrench. Replace any fastener that shows corrosion.", STYLES["BodyText"]),
        Spacer(1, 2 * mm), pdf_table(m["torque"], [70, 30, 40]),
        Paragraph("4. Fault codes", STYLES["Heading2"]),
        Paragraph("The vessel management computer reports faults to the shore operator console.", STYLES["BodyText"]),
        Spacer(1, 2 * mm), pdf_table(m["faults"], [20, 60, 85]),
        Paragraph("5. Battery safety", STYLES["Heading2"]),
    ]
    story += [Paragraph(t, STYLES["BodyText"]) for t in m["battery"]]
    build_pdf(path, story, m["title"])


def docx_doc(title, ref, meta):
    doc = Document()
    doc.core_properties.author = AGENCY
    doc.core_properties.title = title
    doc.core_properties.created = FIXED_TIME
    doc.core_properties.modified = FIXED_TIME
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_paragraph().add_run(BANNER).italic = True
    doc.add_heading(title, level=0)
    doc.add_paragraph(f"{AGENCY}. {ref}. {meta}")
    return doc


def docx_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Light Grid Accent 1"
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            table.cell(r, c).text = str(value)


def write_policy(path, year):
    t = SECURITY[year]
    doc = docx_doc("Security Clearance Policy", t["ref"], t["meta"])
    sections = ["Scope"] + (["Sponsorship"] if t["sponsorship"] else []) + ["Clearance levels", "Validity and renewal",
                                                                          "Security incidents"]
    number = {name: f"{i}. {name}" for i, name in enumerate(sections, 1)}
    doc.add_heading(number["Scope"], level=1)
    doc.add_paragraph("This policy covers security clearances for Blue Defence employees and contractors who work "
                      "with classified material.")
    if t["sponsorship"]:
        doc.add_heading(number["Sponsorship"], level=1)
        doc.add_paragraph(t["sponsorship"])
    doc.add_heading(number["Clearance levels"], level=1)
    for line in t["levels"]:
        doc.add_paragraph(line, style="List Bullet")
    doc.add_heading(number["Validity and renewal"], level=1)
    doc.add_paragraph(t["validity_text"])
    rows = [["Clearance level", "Valid for", "Apply for renewal before expiry"]]
    rows += [[level, valid, renew] for level, (valid, renew) in t["validity"].items()]
    docx_table(doc, rows)
    doc.add_heading(number["Security incidents"], level=1)
    doc.add_paragraph(t["incidents"] + " A security incident includes a lost classified document, an unauthorised "
                      "disclosure and a breach of a restricted area. Preserve any evidence and do not investigate "
                      "the incident yourself.")
    for heading, paras in t["extra"]:
        doc.add_heading(heading, level=1)
        for text in paras:
            doc.add_paragraph(text)
    doc.save(path)


def write_deck(path):
    prs = Presentation()
    prs.core_properties.author = AGENCY
    prs.core_properties.created = FIXED_TIME
    prs.core_properties.modified = FIXED_TIME
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Quarterly Project Status Report, Q3 2026"
    slide.placeholders[1].text = f"{AGENCY} · Programme Management Office · reported 6 October 2026\n{BANNER}"
    for project in PROJECTS:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = project["name"]
        body = slide.placeholders[1].text_frame
        body.text = project["bullets"][0]
        for bullet in project["bullets"][1:]:
            body.add_paragraph().text = bullet
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Portfolio summary"
    rows = [["Project", "Status", "Key date"],
            ["HALCYON", "Amber", "15 December 2026: last 4 vessels"],
            ["LANTERN", "Green", "15 January 2027: go-live"],
            ["ORCHID", "Red", "12 November 2026: steering committee decision"]]
    shape = slide.shapes.add_table(len(rows), 3, Inches(0.5), Inches(1.6), Inches(9), Inches(2)).table
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            shape.cell(r, c).text = value
            shape.cell(r, c).text_frame.paragraphs[0].font.size = PptPt(16)
    prs.save(path)


def write_suppliers(path):
    with path.open("w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(SUPPLIERS)


def write_scanned_memo(path):
    """Render the memo as a page image with skew and noise, so the PDF has no text layer."""
    font_path = os.path.join(os.path.dirname(reportlab.__file__), "fonts", "Vera.ttf")
    width, height = 1654, 2339  # A4 at 200 dpi
    page = Image.new("L", (width, height), 250)
    draw = ImageDraw.Draw(page)
    font = ImageFont.truetype(font_path, 38)
    y = 180
    for line in INSPECTION_MEMO:
        draw.text((170, y), line, fill=25, font=font)
        y += 58
    rng = random.Random(31)
    for _ in range(9000):
        draw.point((rng.randrange(width), rng.randrange(height)), fill=rng.randrange(120, 220))
    page = page.rotate(0.6, resample=Image.BICUBIC, fillcolor=250).filter(ImageFilter.GaussianBlur(0.6))
    page.save(path, "PDF", resolution=200.0)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--out", type=Path, default=Path(__file__).resolve().parent.parent / "sample-docs")
    out = parser.parse_args().out

    kb = out / "blue-defence-knowledge-base"
    for sub in ("", "policies", "archive", "engineering", "reports", "registers"):
        (kb / sub).mkdir(parents=True, exist_ok=True)
    (out / "scanned").mkdir(parents=True, exist_ok=True)

    write_policy(kb / "policies" / "security-clearance-policy-2026.docx", 2026)
    write_policy(kb / "archive" / "security-clearance-policy-2024.docx", 2024)
    write_manual(kb / "engineering" / "sentinel-7-maintenance-manual.pdf")
    write_deck(kb / "reports" / "project-status-q3-2026.pptx")
    write_suppliers(kb / "registers" / "approved-suppliers.csv")
    write_scanned_memo(out / "scanned" / "site-inspection-memo-2026-08-14.pdf")

    for f in sorted(p for p in out.rglob("*") if p.is_file() and p.name != "README.md"):
        print(f"{f.stat().st_size:>8}  {f.relative_to(out)}")


if __name__ == "__main__":
    main()
